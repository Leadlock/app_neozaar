from pathlib import Path
from zipfile import ZIP_STORED, ZipFile, ZipInfo

from lxml import etree
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT


SOURCE = Path("/Users/apple/Downloads/NeoZaar_Services_Managed_IT_Cybersecurity_Proposal_v1.0.docx")
OUTPUT = Path(__file__).resolve().parents[1] / "app" / "proposal-template.docx"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def replace_once(text: str, old: str, new: str) -> str:
    if old not in text:
        raise RuntimeError(f"Missing proposal marker: {old}")
    return text.replace(old, new, 1)


def set_paragraph_text(paragraph, value: str) -> None:
    nodes = paragraph.xpath(".//w:t", namespaces=NS)
    if not nodes:
        raise RuntimeError(f"Cannot rewrite paragraph without text nodes: {value}")
    nodes[0].text = value
    for node in nodes[1:]:
        node.text = ""


REPLACEMENTS = {
    "Basic": "50 Hour Pack",
    "Advanced": "100 Hour Pack",
    "Advanced Plus": "200 Hour Pack",
    "Basic Support": "Reactive Support 50 Hours",
    "Advanced Support": "Reactive Support 100 Hours",
    "Advanced Plus Support": "Reactive Support 200 Hours",
    "Monthly / Coverage Model": "Hour Bank / Validity",
    "Specialist Reserve": "Specialist Access",
    "40 pooled hrs; one customer-local 8h shift x 5 business days": "50 prepaid hrs; valid 3 months; customer-local 8x5 window",
    "Up to 2 L3/SME hours": "L1 through SME; all effort draws from bank",
    "60 pooled hrs; one customer-local 8h shift x 5 business days": "100 prepaid hrs; valid 6 months; customer-local 8x5 window",
    "Up to 5 L3/SME hours": "L1 through SME; all effort draws from bank",
    "100 pooled hrs; SOW-defined support window": "200 prepaid hrs; valid 12 months; SOW-defined window",
    "Up to 10 L3/SME hours": "L1 through SME; all effort draws from bank",
    "Basic Support - 40 pooled support hours/month + up to 2 specialist L3/SME hours; delivered during one customer-local support shift. P1 engineering effort consumes entitlement at 2.0x actual engineer-hours.": "Reactive Support 50 Hours - a prepaid 50-hour bank valid for 3 months and delivered during one customer-local support shift. P1 engineering effort consumes the bank at 2.0x actual engineer-hours.",
    "Advanced Support - 60 pooled support hours/month + up to 5 specialist L3/SME hours; delivered during one customer-local support shift. P1 engineering effort consumes entitlement at 2.0x actual engineer-hours.": "Reactive Support 100 Hours - a prepaid 100-hour bank valid for 6 months and delivered during one customer-local support shift. P1 engineering effort consumes the bank at 2.0x actual engineer-hours.",
    "Advanced Plus Support - 100 pooled support hours/month + up to 10 specialist L3/SME hours; support window is defined in the SOW and may be extended beyond a single shift when commercially agreed.": "Reactive Support 200 Hours - a prepaid 200-hour bank valid for 12 months; its support window is defined in the SOW and may extend beyond a single shift when commercially agreed.",
    "* Basic and Advanced are delivered only during one Customer Local Support Shift (CLSS): one continuous 8-hour support window per business day, five business days per week, aligned to the customer local timezone. Exact start/end times, working days and public-holiday calendar are fixed in the SOW. P1 response clocks for these packages run within that shift unless separate after-hours coverage is expressly contracted. P1 active engineering effort is charged at 2.0x actual engineer-hours. Advanced Plus response targets are measured within its SOW-defined support window.": "* The 50 and 100 Hour Packs are delivered only during one Customer Local Support Shift (CLSS): one continuous 8-hour support window per business day, five business days per week, aligned to the customer local timezone. Exact start/end times, working days and public-holiday calendar are fixed in the SOW. P1 response clocks for these packs run within that shift unless separate after-hours coverage is expressly contracted. P1 active engineering effort is charged at 2.0x actual engineer-hours. The 200 Hour Pack response targets are measured within its SOW-defined support window.",
    "For Advanced and above: cross-domain XDR investigation across endpoint, identity, email, network, cloud and server telemetry where those sources are integrated.": "For the 100 Hour Pack and above: cross-domain XDR investigation across endpoint, identity, email, network, cloud and server telemetry where those sources are integrated.",
    "For Advanced and above: virtual machines, networking, storage, identity/access, security configuration, availability and baseline resource troubleshooting across agreed cloud platforms.": "For the 100 Hour Pack and above: virtual machines, networking, storage, identity/access, security configuration, availability and baseline resource troubleshooting across agreed cloud platforms.",
    "For Advanced and above: security posture reviews, architecture discussions, hardening recommendations, control reviews, remediation planning and technology optimization.": "For the 100 Hour Pack and above: security posture reviews, architecture discussions, hardening recommendations, control reviews, remediation planning and technology optimization.",
    "For Advanced Plus and Enterprise: database availability, connectivity, performance troubleshooting, backup/restore assistance, users/permissions, monitoring, configuration, capacity guidance and health checks for agreed technologies.": "For the 200 Hour Pack and Enterprise: database availability, connectivity, performance troubleshooting, backup/restore assistance, users/permissions, monitoring, configuration, capacity guidance and health checks for agreed technologies.",
    "For Advanced Plus and Enterprise: middleware availability, configuration, integration, application connectivity, log and performance troubleshooting.": "For the 200 Hour Pack and Enterprise: middleware availability, configuration, integration, application connectivity, log and performance troubleshooting.",
    "6.1 Basic Support - 40 Hours/Month": "6.1 Reactive Support 50 Hours",
    "40 pooled support hours per month.": "50 prepaid reactive support hours valid for 3 months from service activation.",
    "Up to 2 L3/SME specialist hours in addition to the pooled support-hour bank.": "L1, L2, L3 and SME effort may be used as required; all actual effort is deducted from the same 50-hour bank.",
    "6.2 Advanced Support - 60 Hours/Month": "6.2 Reactive Support 100 Hours",
    "60 pooled support hours per month.": "100 prepaid reactive support hours valid for 6 months from service activation.",
    "Up to 5 L3/SME specialist hours in addition to the pooled support-hour bank.": "L1, L2, L3 and SME effort may be used as required; all actual effort is deducted from the same 100-hour bank.",
    "Includes Basic capabilities plus XDR, cloud, security consulting, NAC, NDR, CTEM, SIEM, SOAR and UEBA.": "Includes 50 Hour Pack capabilities plus XDR, cloud, security consulting, NAC, NDR, CTEM, SIEM, SOAR and UEBA.",
    "6.3 Advanced Plus Support - 100 Hours/Month": "6.3 Reactive Support 200 Hours",
    "100 pooled support hours per month.": "200 prepaid reactive support hours valid for 12 months from service activation.",
    "Up to 10 L3/SME specialist hours in addition to the pooled support-hour bank.": "L1, L2, L3 and SME effort may be used as required; all actual effort is deducted from the same 200-hour bank.",
    "Includes Advanced capabilities plus DBA, middleware and DevOps specialist support.": "Includes 100 Hour Pack capabilities plus DBA, middleware and DevOps specialist support.",
    "Basic and Advanced operate only within one Customer Local Support Shift (CLSS): one continuous 8-hour support window per business day, five business days per week, aligned to the customer local timezone. The exact start/end times, working days and public-holiday calendar are recorded in the SOW.": "The 50 and 100 Hour Packs operate only within one Customer Local Support Shift (CLSS): one continuous 8-hour support window per business day, five business days per week, aligned to the customer local timezone. The exact start/end times, working days and public-holiday calendar are recorded in the SOW.",
    "Advanced Plus operates within the support window expressly defined in the SOW. Its window may extend beyond a single customer-local shift when commercially agreed; no extended-hours commitment applies unless stated in the SOW.": "The 200 Hour Pack operates within the support window expressly defined in the SOW. Its window may extend beyond a single customer-local shift when commercially agreed; no extended-hours commitment applies unless stated in the SOW.",
    "Tickets may be submitted outside the applicable support window, but for Basic and Advanced the SLA clock and engineering response begin when the Customer Local Support Shift opens unless separate after-hours coverage is expressly contracted. Any agreed after-hours P1 work follows the applicable commercial schedule; unless the SOW states otherwise, the 2.0x P1 consumption rule still applies.": "Tickets may be submitted outside the applicable support window, but for the 50 and 100 Hour Packs the SLA clock and engineering response begin when the Customer Local Support Shift opens unless separate after-hours coverage is expressly contracted. Any agreed after-hours P1 work follows the applicable commercial schedule; unless the SOW states otherwise, the 2.0x P1 consumption rule still applies.",
    "Pooled support hours are consumed based on actual engineer/service effort for active engineering, investigation, administration, change execution, customer-requested technical meetings, and vendor/OEM coordination related to an in-scope ticket.": "Prepaid support hours are consumed from the selected bank based on actual engineer/service effort for active engineering, investigation, administration, change execution, customer-requested technical meetings, and vendor/OEM coordination related to an in-scope ticket.",
    "For Basic and Advanced, all active engineering effort on a P1 incident is subject to a 2.0x support-hour consumption multiplier. One actual engineer-hour therefore consumes two support hours from the applicable entitlement.": "For the 50 and 100 Hour Packs, all active engineering effort on a P1 incident is subject to a 2.0x support-hour consumption multiplier. One actual engineer-hour therefore consumes two support hours from the applicable bank.",
    "L3/SME specialist reserve is additional to the pooled monthly support-hour bank and is used only when specialist escalation is required; for Basic and Advanced P1 work, specialist effort is subject to the 2.0x multiplier unless the SOW states otherwise.": "There is no separate specialist reserve. L3/SME effort is deducted from the selected hour bank; for P1 work under the 50 and 100 Hour Packs, specialist effort is subject to the 2.0x multiplier unless the SOW states otherwise.",
    "Unused monthly pooled hours and specialist reserve expire at month end unless the commercial schedule explicitly allows rollover.": "Unused hours expire at the end of the selected pack validity period. They do not refresh monthly and do not roll over unless a paid validity extension or other written commercial exception is agreed before expiry.",
    "Work exceeding the monthly entitlement requires customer approval and is either deferred to the next period or billed at the agreed overage rate. Where Basic or Advanced P1 work continues beyond available entitlement with approval, the 2.0x consumption rule remains applicable unless the commercial schedule states otherwise.": "Work after the hour bank is exhausted requires customer approval and either a new pack or billing at the agreed overage rate. Where P1 work under the 50 or 100 Hour Pack continues beyond available hours with approval, the 2.0x consumption rule remains applicable unless the commercial schedule states otherwise.",
    "* For Basic and Advanced, response targets are measured only while the Customer Local Support Shift is open. Advanced Plus response targets are measured within its SOW-defined support window. Enterprise targets operate against the contracted 24x7x365 service. The final SLA schedule in the executed SOW supersedes this baseline.": "* For the 50 and 100 Hour Packs, response targets are measured only while the Customer Local Support Shift is open. The 200 Hour Pack response targets are measured within its SOW-defined support window. Enterprise targets operate against the contracted 24x7x365 service. The final SLA schedule in the executed SOW supersedes this baseline.",
    "Monthly service report; review on request or quarterly when agreed.": "Hour-consumption statement; service review on request.",
    "Monthly service report and periodic operational review.": "Hour-consumption statement and periodic operational review.",
    "Monthly service report and technical service review.": "Hour-consumption statement and technical service review.",
    "Monthly hours are service capacity, not a guarantee that all outstanding work can be completed within a given month.": "Purchased hours are a finite service-capacity bank, not a guarantee that all outstanding work can be completed before the bank is exhausted or expires.",
    "Monthly Price": "Pack Price",
    "40 pooled hrs + up to 2 L3/SME; one customer-local 8x5 shift; P1 effort at 2.0x": "50 prepaid hrs; 3-month validity; customer-local 8x5 shift; P1 effort at 2.0x",
    "60 pooled hrs + up to 5 L3/SME; one customer-local 8x5 shift; P1 effort at 2.0x": "100 prepaid hrs; 6-month validity; customer-local 8x5 shift; P1 effort at 2.0x",
    "100 pooled hrs + up to 10 L3/SME; SOW-defined support window": "200 prepaid hrs; 12-month validity; SOW-defined support window",
    "Enterprise Support": "Validity Extension",
    "24x7x365 contracted L1-to-SME model": "Adds one month of validity; does not add support hours",
    "Contract period": "Pack validity",
    "Billing frequency": "Billing model",
    "Basic / Advanced support window": "50 / 100 Hour Pack support window",
    "Basic / Advanced P1 consumption": "50 / 100 Hour Pack P1 consumption",
    "Commercial control: Monthly-hour entitlements are capacity banks, not unlimited-use services. For Basic and Advanced, support is limited to one customer-local shift and P1 active engineering time consumes entitlement at 2.0x actual engineer-hours. Enterprise is likewise governed by the staffing, volume, monitoring and change-capacity assumptions stated in the SOW.": "Commercial control: The selected hour pack is a finite prepaid bank, not an unlimited-use or monthly-refresh service. Unused hours expire at the end of the pack validity period. For the 50 and 100 Hour Packs, support is limited to one customer-local shift and P1 active engineering time consumes the bank at 2.0x actual engineer-hours. Enterprise 24x7 support remains separately scoped and quoted.",
    "1.  Select preferred service package and confirm operating/support window.": "1.  Select the preferred reactive support hour bank and confirm the operating/support window.",
    "Customer Local Support Shift (CLSS): One continuous 8-hour support window per business day, five business days per week, aligned to the customer local timezone. Applicable to Basic and Advanced; exact shift times, working days and holiday calendar are stated in the SOW.": "Customer Local Support Shift (CLSS): One continuous 8-hour support window per business day, five business days per week, aligned to the customer local timezone. Applicable to the 50 and 100 Hour Packs; exact shift times, working days and holiday calendar are stated in the SOW.",
    "Support Hour: Measured engineer/service effort charged against the applicable entitlement. For Basic and Advanced P1 incidents, active actual engineer-hours are multiplied by 2.0x for support-hour consumption.": "Support Hour: Measured engineer/service effort charged against the selected prepaid bank. For P1 incidents under the 50 and 100 Hour Packs, active actual engineer-hours are multiplied by 2.0x for support-hour consumption.",
    "P1 Consumption Multiplier: For Basic and Advanced, a 2.0x factor applied to active actual engineer-hours while an incident is handled as P1. It is an entitlement/billing-consumption rule and does not represent a promise of two engineers or a two-hour elapsed-time commitment.": "P1 Consumption Multiplier: For the 50 and 100 Hour Packs, a 2.0x factor applied to active actual engineer-hours while an incident is handled as P1. It is an hour-bank consumption rule and does not represent a promise of two engineers or a two-hour elapsed-time commitment.",
}


with ZipFile(SOURCE) as source:
    parts = {item.filename: source.read(item.filename) for item in source.infolist()}
    infos = {item.filename: item for item in source.infolist()}

root = etree.fromstring(parts["word/document.xml"])
seen = set()
for paragraph in root.xpath(".//w:p", namespaces=NS):
    text = "".join(paragraph.xpath(".//w:t/text()", namespaces=NS)).strip()
    if text in REPLACEMENTS:
        set_paragraph_text(paragraph, REPLACEMENTS[text])
        seen.add(text)

missing = sorted(set(REPLACEMENTS) - seen)
if missing:
    raise RuntimeError(f"Proposal paragraphs not found: {missing}")

xml = etree.tostring(root, encoding="unicode")
xml = xml.replace("[CUSTOMER NAME]", "Proposal Customer")
xml = replace_once(xml, "23 September 2026", "26 September 2026")
xml = replace_once(xml, "1.1 - Customer Ready", "1.3 - Customer Ready")
xml = replace_once(xml, ">1.1<", ">1.3<")
xml = replace_once(
    xml,
    "Items shown in square brackets must be completed before issue to a customer. The service scope is operationally usable as written, but pricing, customer identity, exact business-hours timezone, supported-technology inventory, and any customer-specific regulatory clauses must be completed in the final SOW.",
    "This proposal is generated from the approved service catalogue. Customer identity, selected support package, commercial model, applicable validity and catalogue pricing are populated by the calculator. Exact timezone, shift hours, holiday calendar, supported-technology inventory, service volumes, customer-specific regulatory clauses, overage rates, onsite rates and any Enterprise 24x7 sizing must be confirmed in the final SOW where applicable.",
)
xml = replace_once(
    xml,
    "Four service levels are available:",
    "Selected package: Calculator Selection. Monthly retainers, prepaid term-valid hour banks and Enterprise 24x7 support are available. The prepaid banks are summarized below; monthly retainers, SLA mapping and optional add-ons are detailed in Appendix C:",
)
xml = replace_once(xml, "[INSERT PRICE]", "INR 279,504 one time")
xml = replace_once(xml, "[INSERT PRICE]", "INR 521,640 one time")
xml = replace_once(xml, "[INSERT PRICE]", "INR 943,920 one time")
xml = replace_once(xml, "[INSERT PRICE / SIZING]", "INR 24,840 per one-month extension")
xml = replace_once(xml, "[12 / 24 / 36 months]", "Calculator Validity")
xml = replace_once(xml, "[Monthly / Quarterly / Annual]", "Prepaid in full")
xml = replace_once(xml, "[INSERT TERMS]", "As agreed in the executed commercial schedule")
xml = replace_once(xml, "[INSERT CURRENCY]", "INR")
xml = replace_once(xml, "[INSERT RATE OR “BY QUOTATION”]", "By quotation")
xml = replace_once(xml, "[INSERT IF APPLICABLE]", "By quotation")
parts["word/document.xml"] = xml.encode("utf-8")

core = parts["docProps/core.xml"].decode("utf-8")
core = core.replace("<dc:title></dc:title>", "<dc:title>Managed IT Infrastructure and Cybersecurity Reactive Support Services</dc:title>")
parts["docProps/core.xml"] = core.encode("utf-8")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
with ZipFile(OUTPUT, "w", compression=ZIP_STORED) as target:
    for name, data in parts.items():
        original = infos[name]
        item = ZipInfo(name, original.date_time)
        item.external_attr = original.external_attr
        item.compress_type = ZIP_STORED
        target.writestr(item, data)


def style_table(table) -> None:
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True


document = Document(OUTPUT)
document.add_page_break()
document.add_heading("Appendix C. Support Commercial Models and Optional Services", level=1)
document.add_paragraph(
    "NeoZaar offers two hour-based commercial models in addition to separately sized Enterprise support. "
    "The selected model, support window, exact technology scope, contract term, rollover exception if any, "
    "and customer-specific operating assumptions must be confirmed in the executed SOW."
)

document.add_heading("Monthly Retainer Packages", level=2)
document.add_paragraph(
    "Monthly retainer hours replenish at the start of each monthly service period. Unused monthly hours expire "
    "at month end unless the executed commercial schedule expressly permits rollover."
)
retainers = document.add_table(rows=1, cols=4)
style_table(retainers)
for cell, value in zip(retainers.rows[0].cells, ["Package", "Monthly Entitlement", "Service Model", "Monthly Price"]):
    cell.text = value
for values in [
    ("40 Hour Monthly Retainer", "40 pooled hours plus up to 2 L3 or SME hours", "Customer-local 8x5 shift; P1 effort at 2.0x", "INR 186,336 per month"),
    ("60 Hour Monthly Retainer", "60 pooled hours plus up to 5 L3 or SME hours", "Customer-local 8x5 shift; P1 effort at 2.0x", "INR 260,820 per month"),
    ("100 Hour Monthly Retainer", "100 pooled hours plus up to 10 L3 or SME hours", "SOW-defined support window; P1 effort at 1.0x", "INR 393,300 per month"),
]:
    cells = retainers.add_row().cells
    for cell, value in zip(cells, values):
        cell.text = value

document.add_heading("Prepaid Term Valid Hour Banks", level=2)
document.add_paragraph(
    "The 50, 100 and 200-hour prepaid banks remain available with 3, 6 and 12-month validity respectively. "
    "They do not replenish monthly. Unused hours expire at the end of the validity period unless a paid validity "
    "extension or another written commercial exception is agreed before expiry. Pricing appears in Section 22."
)

document.add_heading("SLA Mapping Across Commercial Models", level=2)
sla = document.add_table(rows=1, cols=5)
style_table(sla)
for cell, value in zip(sla.rows[0].cells, ["Package", "P1", "P2", "P3", "P4"]):
    cell.text = value
for values in [
    ("40 Hour Monthly Retainer or 50 Hour Bank", "60 min", "4 business hrs", "8 business hrs", "2 business days"),
    ("60 Hour Monthly Retainer or 100 Hour Bank", "60 min", "2 business hrs", "4 business hrs", "1 business day"),
    ("100 Hour Monthly Retainer or 200 Hour Bank", "30 min", "1 business hr", "4 business hrs", "1 business day"),
    ("Enterprise 24x7x365", "15 min", "30 min", "2 hrs", "4 hrs"),
]:
    cells = sla.add_row().cells
    for cell, value in zip(cells, values):
        cell.text = value
document.add_paragraph(
    "Response targets remain subject to the applicable service window and the final SLA schedule in the executed SOW."
)

document.add_page_break()
document.add_heading("Optional Additional Services", level=2)
document.add_paragraph(
    "Selected optional services: Calculator Add-ons. All add-ons are separately scoped and priced by quotation."
)
addons = document.add_table(rows=1, cols=2)
style_table(addons)
addons.rows[0].cells[0].text = "Assessment and Advisory"
addons.rows[0].cells[1].text = "Security, Transformation and Delivery"
left = [
    "Cybersecurity maturity assessment", "Infrastructure assessment", "Cloud security assessment",
    "Active Directory security assessment", "Firewall rule review", "Microsoft 365 security assessment",
    "Backup and disaster-recovery assessment", "Architecture consulting", "Zero Trust consulting",
    "Security awareness program",
]
right = [
    "Incident response retainer", "Vulnerability management", "Penetration testing", "Red-team exercise",
    "Cloud migration", "Infrastructure modernization", "SOC implementation", "SIEM implementation and tuning",
    "Dedicated engineering or onsite services", "NeoZaar CloudCraft deployment services",
]
for a, b in zip(left, right):
    cells = addons.add_row().cells
    cells[0].text = a
    cells[1].text = b

working = OUTPUT.with_suffix(".working.docx")
document.save(working)
with ZipFile(working) as source:
    updated_parts = {item.filename: source.read(item.filename) for item in source.infolist()}
    updated_infos = {item.filename: item for item in source.infolist()}
with ZipFile(OUTPUT, "w", compression=ZIP_STORED) as target:
    for name, data in updated_parts.items():
        original = updated_infos[name]
        item = ZipInfo(name, original.date_time)
        item.external_attr = original.external_attr
        item.compress_type = ZIP_STORED
        target.writestr(item, data)
working.unlink()

print(OUTPUT)
